"""
generate_report.py
Creates reports/Supermarket_Sales_Report.docx with full analysis insights.
Run once:  python generate_report.py
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime
import os

from data_cleaning import load_and_clean
from analysis import run_analysis


# ── helpers ───────────────────────────────────────────────────────────────────

def add_heading(doc, text, level=1, color=(30, 41, 59)):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        run.font.color.rgb = RGBColor(*color)
    return h


def add_paragraph(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p


def shade_row(row, hex_color="D6E4F7"):
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)


def add_table(doc, headers, rows, header_color="2563EB"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    hdr = table.rows[0]
    shade_row(hdr, "2563EB")
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        if ri % 2 == 0:
            shade_row(row, "EFF6FF")
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = str(val)
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)

    return table


# ── main ──────────────────────────────────────────────────────────────────────

def generate(output_path="reports/Supermarket_Sales_Report.docx"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    df = load_and_clean()
    results = run_analysis(df)

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)

    # ── Cover ────────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Supermarket Sales Analysis")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(30, 41, 59)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Business Insights Report")
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(100, 116, 139)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = date_p.add_run(f"Generated: {datetime.date.today().strftime('%B %d, %Y')}")
    r2.font.size = Pt(11)
    r2.italic = True
    r2.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_paragraph()

    # ── 1. Project Overview ──────────────────────────────────────────────
    add_heading(doc, "1. Project Overview", level=1)
    add_paragraph(
        doc,
        "This report summarises the analysis of 500 supermarket transactions recorded across four "
        "Indian cities (Jaipur, Delhi, Mumbai, Bengaluru) between January and July 2026. The goal "
        "is to derive actionable business insights covering product performance, branch revenue, "
        "customer behaviour, and payment preferences.",
    )

    # ── 2. Dataset Summary ───────────────────────────────────────────────
    add_heading(doc, "2. Dataset Summary", level=1)
    add_table(
        doc,
        headers=["Field", "Details"],
        rows=[
            ["File", "Supermarket-data.csv"],
            ["Total Rows", "500 transactions"],
            ["Columns", "13 (Invoice ID, Date, Branch, City, Customer Type, Gender, Product, Category, Quantity, Unit Price, Payment, Rating, Sales)"],
            ["Branches", "A – Jaipur, B – Delhi, C – Mumbai, D – Bengaluru"],
            ["Date Range", "Jan 2026 – Jul 2026"],
            ["Categories", "Beverages, Personal Care, Dairy, Grocery, Fruits, Snacks, Vegetables, Bakery"],
            ["Missing Values", "None"],
            ["Duplicate Rows", "None"],
        ],
    )
    doc.add_paragraph()

    # ── 3. Key Performance Indicators ───────────────────────────────────
    add_heading(doc, "3. Key Performance Indicators", level=1)
    total_rev = df["Sales"].sum()
    avg_order = df["Sales"].mean()
    avg_rating = df["Rating"].mean()
    member_pct = (df["Customer Type"] == "Member").mean() * 100

    add_table(
        doc,
        headers=["KPI", "Value"],
        rows=[
            ["Total Revenue", f"Rs {total_rev:,.2f}"],
            ["Total Transactions", f"{len(df):,}"],
            ["Average Order Value", f"Rs {avg_order:,.2f}"],
            ["Average Customer Rating", f"{avg_rating:.2f} / 5"],
            ["Member Customers", f"{member_pct:.1f}% of all transactions"],
            ["Cities Covered", "4 (Jaipur, Delhi, Mumbai, Bengaluru)"],
        ],
    )
    doc.add_paragraph()

    # ── 4. Business Insights ──────────────────────────────────────────────
    add_heading(doc, "4. Business Insights", level=1)

    # 4.1 Top Product
    add_heading(doc, "4.1  Top Products by Sales", level=2)
    add_paragraph(
        doc,
        f"Cheese leads all products with Rs {results['top_product_sales']:,.2f} in total sales. "
        "Coffee and Shampoo follow closely, suggesting strong demand in the Beverages and Personal "
        "Care segments. Basic staples like Biscuits and Milk trail behind premium products.",
    )
    prod_rows = [
        [row["Product"], f"Rs {row['Total Sales']:,.2f}"]
        for _, row in results["product_sales"].iterrows()
    ]
    add_table(doc, ["Product", "Total Sales"], prod_rows)
    doc.add_paragraph()

    # 4.2 Branch Performance
    add_heading(doc, "4.2  Branch Performance", level=2)
    add_paragraph(
        doc,
        f"Branch C (Mumbai) is the highest-performing branch with Rs {results['best_branch_sales']:,.2f} "
        "in revenue, ahead of Branch B (Delhi). Branch A (Jaipur) records the lowest revenue — "
        "targeted promotions there could unlock additional growth.",
    )
    branch_rows = [
        [row["Branch_Label"], f"Rs {row['Total Sales']:,.2f}"]
        for _, row in results["branch_sales"].iterrows()
    ]
    add_table(doc, ["Branch", "Total Sales"], branch_rows)
    doc.add_paragraph()

    # 4.3 Category Breakdown
    add_heading(doc, "4.3  Category Breakdown", level=2)
    add_paragraph(
        doc,
        f"Beverages is the top-grossing category (Rs {results['top_category_sales']:,.2f}), "
        "driven by high unit prices for Coffee and Tea. Personal Care and Dairy follow. "
        "Bakery records the lowest category revenue.",
    )
    cat_rows = [
        [row["Category"], f"Rs {row['Total Sales']:,.2f}"]
        for _, row in results["category_sales"].iterrows()
    ]
    add_table(doc, ["Category", "Total Sales"], cat_rows)
    doc.add_paragraph()

    # 4.4 Payment Method
    add_heading(doc, "4.4  Payment Method Popularity", level=2)
    add_paragraph(
        doc,
        f"Payment methods are almost equally distributed. UPI leads with {results['top_payment_count']} "
        "transactions, closely followed by Net Banking (126), Card (125), and Cash (122). "
        "This indicates a highly digital-savvy customer base.",
    )
    pay_rows = [
        [row["Payment"], str(row["Count"])]
        for _, row in results["payment_counts"].iterrows()
    ]
    add_table(doc, ["Payment Method", "Transactions"], pay_rows)
    doc.add_paragraph()

    # 4.5 Member vs Normal
    add_heading(doc, "4.5  Member vs Normal Customer Spending", level=2)
    add_paragraph(
        doc,
        "Member customers account for 59.2% of all transactions and generate 58.5% of total "
        "revenue (Rs 1,43,009 vs Rs 1,01,402 for Normal customers). However, the average order "
        "value is slightly higher for Normal customers (Rs 497 vs Rs 483), suggesting Members "
        "transact more frequently.",
    )
    cust_rows = [
        [
            row["Customer Type"],
            f"Rs {row['Total']:,.2f}",
            f"Rs {row['Average']:,.2f}",
            str(int(row["Transactions"])),
        ]
        for _, row in results["customer_spend"].iterrows()
    ]
    add_table(doc, ["Customer Type", "Total Spend", "Avg Order", "Transactions"], cust_rows)
    doc.add_paragraph()

    # 4.6 Ratings
    add_heading(doc, "4.6  Customer Ratings", level=2)
    add_paragraph(
        doc,
        f"The overall average rating is {results['avg_rating']} / 5. Branch D (Bengaluru) "
        "scores highest at 4.09 while Branch A (Jaipur) scores lowest at 3.84. Across all "
        "branches ratings cluster between 3.0 and 5.0, with the median near 4.0.",
    )
    rating_rows = [
        [f"Branch {row['Branch']}", str(row["Avg Rating"])]
        for _, row in results["rating_by_branch"].iterrows()
    ]
    add_table(doc, ["Branch", "Avg Rating"], rating_rows)
    doc.add_paragraph()

    # ── 5. Dashboard Layout ────────────────────────────────────────────────
    add_heading(doc, "5. Streamlit Dashboard — Layout Description", level=1)
    add_paragraph(
        doc,
        "The interactive dashboard (app.py) is structured into the following sections:",
        bold=True,
    )
    sections = [
        ("Sidebar", "Multi-select filters for Branch, Category, Customer Type, Payment Method, and a Date Range picker. All charts react to filter changes in real time."),
        ("KPI Row", "Six metric cards: Total Revenue, Transactions, Avg Order Value, Avg Rating, Member count, and Cities served."),
        ("Product & Branch", "Side-by-side: horizontal bar chart (Products by Sales) and grouped bar chart (Branch Revenue)."),
        ("Category & Payment", "Donut pie chart (Category sales share) and bar chart (Payment transaction counts)."),
        ("Monthly Trend", "Spline line chart showing revenue month-over-month from Jan to Jul 2026."),
        ("Customer Insights", "Three columns: Member vs Normal total spend, Gender sales split (pie), and Avg Rating by Branch."),
        ("Heatmap", "Sales matrix heat-map crossing Category (rows) against Branch (columns)."),
        ("Raw Data Table", "Collapsible, filterable, scrollable table of all transactions."),
    ]
    for title, desc in sections:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{title}: ")
        run.bold = True
        run.font.size = Pt(10.5)
        p.add_run(desc).font.size = Pt(10.5)

    doc.add_paragraph()

    # ── 6. Recommendations ────────────────────────────────────────────────
    add_heading(doc, "6. Recommendations", level=1)
    recs = [
        "Expand the Beverages and Personal Care range in Branch A (Jaipur) to close the revenue gap with Mumbai.",
        "Launch a loyalty bonus for Normal customers to encourage conversion to membership.",
        "Offer cashback on Card payments to balance the slight lag behind UPI and Net Banking.",
        "Investigate the lower ratings at Branch A and B through customer feedback surveys.",
        "Introduce seasonal promotions on Bakery items, which lag all other categories.",
    ]
    for rec in recs:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(rec).font.size = Pt(10.5)

    doc.add_paragraph()

    # ── Footer note ────────────────────────────────────────────────────────
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_p.add_run("Supermarket Sales Analysis  |  Generated with Python & python-docx  |  2026")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(148, 163, 184)
    r.italic = True

    doc.save(output_path)
    print(f"[OK] Report saved -> {output_path}")


if __name__ == "__main__":
    generate()
